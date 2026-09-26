"""JSONL Jupyter bridge. Kernels execute on the host, NOT in a security sandbox."""
from __future__ import annotations

import argparse
import asyncio
import contextlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import platform
import re
import shutil
import sys
import threading
import time
import uuid
from datetime import datetime, timezone

MAX_OUTPUT_BYTES = 8 * 1024 * 1024
MAX_OUTPUT_MESSAGES = 500
MAX_MIME_BYTES = 4 * 1024 * 1024
SECRET_KEY = re.compile(r"TOKEN|SECRET|PASSWORD|PASSWD|CREDENTIAL|API.?KEY|PRIVATE.?KEY|AUTH|COOKIE", re.I)
PACKAGE_SPEC = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]*(?:\[[A-Za-z0-9_,.-]+\])?(?:(?:===|==|!=|~=|>=|<=|>|<)[A-Za-z0-9.*+!_-]+(?:,(?:==|!=|~=|>=|<=|>|<)[A-Za-z0-9.*+!_-]+)*)?$")


def now():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def emit(value):
    def finite(item):
        if isinstance(item, float) and not math.isfinite(item):
            return None
        if isinstance(item, dict):
            return {str(k): finite(v) for k, v in item.items()}
        if isinstance(item, (list, tuple)):
            return [finite(v) for v in item]
        return item
    print(json.dumps(finite(value), ensure_ascii=False, default=str, allow_nan=False), flush=True)


def clean_env():
    # Keep ordinary platform variables, but never pass inherited credentials or
    # Python startup hooks to the worker's kernels/package subprocesses.
    env = {k: v for k, v in os.environ.items() if not SECRET_KEY.search(k)}
    for key in ("PYTHONPATH", "PYTHONSTARTUP", "PYTHONHOME", "NODE_OPTIONS", "LD_PRELOAD"):
        env.pop(key, None)
    env.update(PYTHONUTF8="1", PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1", PYTHONNOUSERSITE="1")
    return env


def python_at(folder):
    return str(folder / ("Scripts/python.exe" if os.name == "nt" else "bin/python"))


class Bridge:
    def __init__(self, root, idle_seconds=1800):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.kernels = {}
        self.tasks = set()
        self.stopping = False
        self.idle_seconds = idle_seconds
        self.runtime_root = Path(os.environ.get("SCIENCE_RUNTIME_DIR") or os.environ.get("ROOT_RUNTIME")
                                 or Path(__file__).resolve().parent.parent / ".runtime").resolve()
        self.environment_root = Path(os.environ.get("SCIENCE_ENVIRONMENT_DIR")
                                     or self.root / ".science-environments").resolve()
        self.environment_locks = {}

    def local_r(self):
        runtime = self.runtime_root / "R"
        executable = Path(os.environ.get("SCIENCE_R_PATH") or runtime / "app" / "bin" / "x64" / "R.exe").resolve()
        rscript_options = [executable.parent / "Rscript.exe", executable.parent.parent / "Rscript.exe",
                           runtime / "app" / "bin" / "Rscript.exe"]
        rscript = next((item for item in rscript_options if item.is_file()), rscript_options[0])
        library = runtime / "library"
        return {"path": str(runtime / "app"), "executable": str(executable),
                "rscript": str(rscript), "library": str(library),
                "available": executable.is_file() and (library / "IRkernel" / "DESCRIPTION").is_file()}

    def r_env(self, library):
        env = clean_env()
        for key in list(env):
            if key == "LANG" or key.startswith("LC_") or key in ("R_HOME", "R_PROFILE", "R_ENVIRON", "R_PROFILE_USER", "R_ENVIRON_USER"):
                env.pop(key, None)
        env.update(R_LIBS_USER=library, R_LIBS_SITE=library, LANGUAGE="en")
        return env

    def contained(self, value, must_exist=True):
        candidate = Path(value)
        candidate = (self.root / candidate).resolve() if not candidate.is_absolute() else candidate.resolve()
        if not candidate.is_relative_to(self.root):
            raise ValueError("Path must be inside the configured project root")
        if must_exist and not candidate.exists():
            raise ValueError("Path does not exist")
        return candidate

    def event(self, kind, **value):
        emit({"event": kind, **value})

    def public_kernel(self, kernel):
        return {key: kernel.get(key) for key in (
            "id", "language", "status", "projectId", "sessionId", "cwd", "createdAt",
            "lastActivity", "environmentId", "environment", "resources", "limits"
        )}

    def get_kernel(self, params):
        kid = params.get("kernelId") or params.get("id")
        if kid not in self.kernels:
            raise ValueError("Kernel does not exist")
        return self.kernels[kid]

    def status(self, kernel, value, execution_id=None):
        kernel["status"] = value
        kernel["lastActivity"] = now()
        kernel["lastMonotonic"] = time.monotonic()
        self.event("status", kernelId=kernel["id"], executionId=execution_id, status=value)

    async def probe(self, params=None):
        dependencies = {}
        for name in ("jupyter_client", "ipykernel", "matplotlib", "numpy", "pandas", "psutil", "pyarrow", "openpyxl"):
            try:
                dependencies[name] = importlib.metadata.version(name)
            except importlib.metadata.PackageNotFoundError:
                dependencies[name] = None
        specs = []
        local_r = self.local_r()
        r_available = local_r["available"]
        if r_available:
            specs.append({"name": "science-r", "language": "r", "displayName": "R (app-local)", "available": True})
        if dependencies["jupyter_client"]:
            from jupyter_client.kernelspec import KernelSpecManager
            for name, entry in KernelSpecManager().get_all_specs().items():
                spec = entry["spec"]
                argv = spec.get("argv", [])
                executable = argv[0] if argv else ""
                available = bool(executable and (Path(executable).is_file() or shutil.which(executable)))
                specs.append({"name": name, "language": spec.get("language"), "displayName": spec.get("display_name"), "available": available})
                r_available |= spec.get("language", "").lower() == "r" and available
        return {"python": {"available": bool(dependencies["jupyter_client"] and dependencies["ipykernel"]),
                           "path": sys.executable, "version": platform.python_version()},
                "r": {"available": r_available, "path": local_r["executable"] if local_r["available"] else None,
                      "reason": None if r_available else "R with a registered IRkernel is not installed"},
                "dependencies": dependencies, "kernelspecs": specs, "executionMode": "host", "sandboxed": False}

    def environment(self, environment_id="default"):
        if environment_id == "default":
            prefix = Path(sys.prefix).resolve()
            managed = prefix.is_relative_to(self.runtime_root / "python") or prefix == Path(__file__).resolve().parent.parent / ".venv"
            return {"id": "default", "name": "Scientific Python", "language": "python", "path": sys.prefix,
                    "pythonPath": sys.executable, "managed": managed}
        if environment_id == "r-default":
            runtime = self.local_r()
            if not runtime["available"]:
                raise ValueError("App-local R runtime is not installed")
            return {"id": "r-default", "name": "Scientific R", "language": "r", "path": runtime["library"],
                    "rPath": runtime["executable"], "managed": True}
        if not re.fullmatch(r"[a-f0-9-]{36}", environment_id):
            raise ValueError("Unknown environment")
        folder = self.environment_root / environment_id
        metadata = folder / "science-environment.json"
        if not metadata.is_file():
            raise ValueError("Unknown environment")
        result = json.loads(metadata.read_text(encoding="utf-8"))
        result.update(id=environment_id, path=str(folder), managed=True)
        if result.get("language") == "r":
            result["rPath"] = self.local_r()["executable"]
        else:
            result["pythonPath"] = python_at(folder)
        return result

    async def launch(self, kernel):
        from jupyter_client import AsyncKernelManager
        from jupyter_client.kernelspec import KernelSpec
        language = kernel["language"]
        kernel.pop("process", None)
        kernel.pop("processCache", None)
        env = clean_env()
        if language == "python":
            environment = self.environment(kernel["environmentId"])
            if environment.get("language") == "r":
                raise ValueError("Choose a Python environment for a Python kernel")
            executable = environment["pythonPath"]
            manager = AsyncKernelManager(kernel_name="python3", autorestart=False)
            # Never use a user python3 kernelspec: launch exactly the selected env.
            manager._kernel_spec = KernelSpec(argv=[executable, "-m", "ipykernel_launcher", "-f", "{connection_file}"],
                                             display_name="Scientific Python", language="python", env={})
            env["PATH"] = str(Path(executable).parent) + os.pathsep + env.get("PATH", "")
            env["VIRTUAL_ENV"] = environment["path"]
        else:
            if self.local_r()["available"]:
                environment = self.environment(kernel["environmentId"])
                if environment.get("language") != "r":
                    raise ValueError("Choose an R environment for an R kernel")
                env = self.r_env(environment["path"])
                manager = AsyncKernelManager(kernel_name="science-r", autorestart=False)
                manager._kernel_spec = KernelSpec(argv=[environment["rPath"], "--slave", "--vanilla", "-e", "IRkernel::main()",
                                                        "--args", "{connection_file}"],
                    display_name="Scientific R", language="R", env={})
            else:
                if os.environ.get("SCIENCE_R_PATH"):
                    raise ValueError("Configured R executable or app-local IRkernel library is unavailable")
                from jupyter_client.kernelspec import KernelSpecManager
                candidates = [(n, s) for n, s in KernelSpecManager().get_all_specs().items()
                              if s["spec"].get("language", "").lower() == "r"]
                if not candidates:
                    raise ValueError("R is unavailable: install R and register IRkernel first")
                manager = AsyncKernelManager(kernel_name=candidates[0][0], autorestart=False)
                manager.kernel_spec.env = {k: v for k, v in manager.kernel_spec.env.items() if not SECRET_KEY.search(k)}
        kernel["manager"] = manager
        try:
            await manager.start_kernel(cwd=kernel["cwd"], env=env, stdout=sys.stderr, stderr=sys.stderr)
            client = manager.client()
            kernel["client"] = client
            client.start_channels()
            await client.wait_for_ready(timeout=40)
            kernel["environment"] = {"language": language, "cwd": kernel["cwd"], "host": platform.platform(),
                                     "executionMode": "host", "sandboxed": False, "capturedAt": now()}
            if language == "python":
                setup = """import json as _science_json, sys as _science_sys, importlib.metadata as _science_meta
get_ipython().run_line_magic('matplotlib', 'inline')
print(_science_json.dumps({'python': _science_sys.version, 'executable': _science_sys.executable,
 'packages': {d.metadata['Name']: d.version for d in _science_meta.distributions() if d.metadata['Name']}}))
del _science_json, _science_sys, _science_meta
"""
                msg_id = client.execute(setup, silent=False, store_history=False, allow_stdin=False)
                context = {"outputs": [], "bytes": 0, "truncated": False, "count": None, "error": False}
                await asyncio.wait_for(self.collect(kernel, msg_id, None, context, publish=False), 30)
                if context["error"]:
                    raise RuntimeError("Scientific environment initialization failed: " + " ".join(o.get("text", "") for o in context["outputs"]))
                streams = "".join(o.get("text", "") for o in context["outputs"] if o["type"] == "stream")
                for line in streams.splitlines():
                    try:
                        info = json.loads(line)
                        if isinstance(info, dict) and "executable" in info:
                            kernel["environment"].update(info)
                    except json.JSONDecodeError:
                        pass
            else:
                kernel["environment"]["kernelspec"] = manager.kernel_name
                setup = "cat(jsonlite::toJSON(list(r=R.version.string, packages=as.list(setNames(installed.packages()[,'Version'], installed.packages()[,'Package']))),auto_unbox=TRUE), '\\n')"
                msg_id = client.execute(setup, silent=False, store_history=False, allow_stdin=False)
                context = {"outputs": [], "bytes": 0, "truncated": False, "count": None, "error": False}
                await asyncio.wait_for(self.collect(kernel, msg_id, None, context, publish=False), 30)
                for output in context["outputs"]:
                    if output["type"] == "stream":
                        with contextlib.suppress(ValueError, json.JSONDecodeError):
                            kernel["environment"].update(json.loads(output.get("text", "")))
        except BaseException:
            with contextlib.suppress(Exception):
                await manager.shutdown_kernel(now=True)
            if kernel.get("client"):
                kernel["client"].stop_channels()
            raise

    async def create(self, params):
        language = str(params.get("language", "python")).lower()
        if language not in ("python", "r"):
            raise ValueError("Language must be python or r")
        cwd = self.contained(params.get("cwd") or self.root)
        if not cwd.is_dir():
            raise ValueError("Working directory must be a directory")
        max_memory = int(params.get("maxMemoryBytes", 4 * 1024 ** 3))
        if not 64 * 1024 ** 2 <= max_memory <= 256 * 1024 ** 3:
            raise ValueError("Kernel memory limit must be between 64 MiB and 256 GiB")
        kernel = {"id": str(uuid.uuid4()), "language": language, "status": "starting", "cwd": str(cwd),
                  "projectId": params.get("projectId"), "sessionId": params.get("sessionId"),
                  "createdAt": now(), "lastActivity": now(), "lastMonotonic": time.monotonic(),
                  "environmentId": params.get("environmentId") or ("default" if language == "python" else "r-default"), "lock": asyncio.Lock(),
                  "interrupted": False, "closing": False, "activeExecutionId": None, "resources": {},
                  "limits": {"maxMemoryBytes": max_memory, "memoryEnforcement": "sampled-watchdog"}, "resourceLimit": None}
        await self.launch(kernel)
        self.kernels[kernel["id"]] = kernel
        self.status(kernel, "idle")
        return self.public_kernel(kernel)

    def append_output(self, kernel, execution_id, context, output, publish):
        size = len(json.dumps(output, ensure_ascii=False, default=str).encode("utf-8"))
        if context["truncated"]:
            return
        if len(context["outputs"]) >= MAX_OUTPUT_MESSAGES or context["bytes"] + size > MAX_OUTPUT_BYTES:
            context["truncated"] = True
            output = {"type": "stream", "name": "stderr", "text": "\n[Output limit reached. Further output omitted.]\n", "truncated": True}
            size = len(output["text"])
        context["outputs"].append(output)
        context["bytes"] += size
        if publish:
            self.event("output", kernelId=kernel["id"], executionId=execution_id, output=output)

    async def collect(self, kernel, msg_id, execution_id, context, publish=True):
        client = kernel["client"]
        while True:
            msg = await client.get_iopub_msg(timeout=None)
            if msg.get("parent_header", {}).get("msg_id") != msg_id:
                continue
            kind, content = msg.get("msg_type"), msg.get("content", {})
            output = None
            if kind == "status" and content.get("execution_state") == "idle":
                break
            if kind == "stream":
                text = content.get("text", "")
                chunk_truncated = len(text) > 128 * 1024
                if chunk_truncated:
                    text = text[:128 * 1024] + "\n[Stream chunk truncated]\n"
                output = {"type": "stream", "name": content.get("name", "stdout"), "text": text}
                if chunk_truncated:
                    output["truncated"] = True
                    context["partiallyTruncated"] = True
            elif kind in ("display_data", "execute_result", "update_display_data"):
                data = content.get("data", {})
                if len(json.dumps(data, default=str).encode("utf-8")) > MAX_MIME_BYTES:
                    data = {"text/plain": "[Display exceeded 4 MiB. Save the full result to a project file.]"}
                    context["partiallyTruncated"] = True
                output = {"type": "display", "data": data, "metadata": content.get("metadata", {}),
                          "displayId": content.get("transient", {}).get("display_id"), "update": kind == "update_display_data"}
                if "execution_count" in content:
                    context["count"] = content["execution_count"]
            elif kind == "error":
                context["error"] = True
                traceback = content.get("traceback", [])
                output = {"type": "error", "ename": content.get("ename"), "evalue": content.get("evalue"),
                          "traceback": traceback, "text": "\n".join(traceback)}
            elif kind == "clear_output":
                # A clear marker lets consumers replace previous display content.
                output = {"type": "display", "data": {}, "clear": True, "wait": bool(content.get("wait"))}
            elif kind == "execute_input":
                context["count"] = content.get("execution_count")
            if output:
                self.append_output(kernel, execution_id, context, output, publish)
        while True:
            reply = await client.get_shell_msg(timeout=None)
            if reply.get("parent_header", {}).get("msg_id") == msg_id:
                content = reply.get("content", {})
                context["count"] = content.get("execution_count", context["count"])
                context["error"] |= content.get("status") == "error"
                return

    async def execute(self, params):
        kernel = self.get_kernel(params)
        code = params.get("code")
        if not isinstance(code, str) or not code.strip():
            raise ValueError("Code must be a nonempty string")
        if len(code.encode("utf-8")) > 1024 * 1024:
            raise ValueError("Code exceeds 1 MiB")
        timeout = float(params.get("timeout", 120))
        if not 0.1 <= timeout <= 3600:
            raise ValueError("Timeout must be between 0.1 and 3600 seconds")
        execution_id = params.get("executionId") or str(uuid.uuid4())
        async with kernel["lock"]:
            if kernel["closing"] or kernel["status"] == "dead":
                raise ValueError("Kernel is closed or dead; create or restart it")
            started = now()
            kernel["interrupted"] = False
            kernel["resourceLimit"] = None
            kernel["activeExecutionId"] = execution_id
            self.status(kernel, "busy", execution_id)
            context = {"outputs": [], "bytes": 0, "truncated": False, "count": None, "error": False}
            result_status = "completed"
            collector = None
            try:
                msg_id = kernel["client"].execute(code, store_history=True, allow_stdin=False, stop_on_error=True)
                collector = asyncio.create_task(self.collect(kernel, msg_id, execution_id, context))
                kernel["collector"] = collector
                # Shield so timeout does not cancel a receive that just consumed idle.
                await asyncio.wait_for(asyncio.shield(collector), timeout)
                result_status = "interrupted" if kernel["interrupted"] else ("error" if context["error"] else "completed")
            except asyncio.TimeoutError:
                result_status = "timeout"
                with contextlib.suppress(Exception):
                    await kernel["manager"].interrupt_kernel()
                try:
                    await asyncio.wait_for(asyncio.shield(collector), 4)
                except Exception:
                    if collector:
                        collector.cancel()
                    await self.stop_manager(kernel)
                    self.status(kernel, "dead", execution_id)
                self.append_output(kernel, execution_id, context, {"type": "error", "ename": "ExecutionTimeout",
                    "evalue": f"Execution exceeded {timeout:g} seconds", "traceback": [],
                    "text": f"Execution exceeded {timeout:g} seconds; " +
                            ("kernel was stopped after it did not respond to interruption. Restart it to continue."
                             if kernel["status"] == "dead" else "kernel was interrupted.")}, True)
            except asyncio.CancelledError:
                if kernel["resourceLimit"]:
                    result_status = "error"
                    self.append_output(kernel, execution_id, context, {"type": "error", "ename": "ResourceLimitExceeded",
                        "evalue": kernel["resourceLimit"], "traceback": [], "text": kernel["resourceLimit"]}, True)
                elif kernel["interrupted"] or kernel["closing"]:
                    result_status = "interrupted"
                else:
                    raise
            except Exception as exc:
                result_status = "interrupted" if kernel["interrupted"] else "error"
                if collector:
                    collector.cancel()
                self.append_output(kernel, execution_id, context, {"type": "error", "ename": type(exc).__name__,
                    "evalue": str(exc), "traceback": [], "text": str(exc)}, True)
                if not await kernel["manager"].is_alive():
                    self.status(kernel, "dead", execution_id)
            finally:
                if collector and not collector.done():
                    collector.cancel()
                if collector:
                    with contextlib.suppress(BaseException):
                        await collector
                if kernel["resourceLimit"]:
                    result_status = "error"
                    if not any(o.get("ename") == "ResourceLimitExceeded" for o in context["outputs"]):
                        self.append_output(kernel, execution_id, context, {"type": "error", "ename": "ResourceLimitExceeded",
                            "evalue": kernel["resourceLimit"], "traceback": [], "text": kernel["resourceLimit"]}, True)
                kernel["activeExecutionId"] = None
                kernel["collector"] = None
                if kernel["status"] not in ("dead", "closed"):
                    self.status(kernel, "idle", execution_id)
            return {"kernelId": kernel["id"], "executionId": execution_id, "outputs": context["outputs"],
                    "status": result_status, "executionCount": context["count"], "environment": kernel["environment"],
                    "startedAt": started, "endedAt": now(), "truncated": context["truncated"] or context.get("partiallyTruncated", False)}

    async def interrupt(self, params):
        kernel = self.get_kernel(params)
        if kernel["activeExecutionId"]:
            kernel["interrupted"] = True
            await kernel["manager"].interrupt_kernel()
            execution_id = kernel["activeExecutionId"]
            async def enforce_interrupt():
                await asyncio.sleep(4)
                if kernel["activeExecutionId"] == execution_id and kernel["status"] == "busy":
                    await self.stop_manager(kernel)
                    self.status(kernel, "dead", execution_id)
                    if kernel.get("collector"):
                        kernel["collector"].cancel()
            task = asyncio.create_task(enforce_interrupt())
            self.tasks.add(task)
            task.add_done_callback(self.tasks.discard)
        return {"ok": True, "kernelId": kernel["id"], "interrupted": bool(kernel["activeExecutionId"])}

    async def stop_manager(self, kernel):
        # Python's Windows venv launcher can have a child interpreter. Remember
        # only this owned kernel's descendants before stopping the launcher.
        descendants = []
        with contextlib.suppress(Exception):
            import psutil
            descendants = psutil.Process(kernel["manager"].provisioner.pid).children(recursive=True)
        with contextlib.suppress(Exception):
            await asyncio.wait_for(kernel["manager"].shutdown_kernel(now=True), 8)
        for child in reversed(descendants):
            with contextlib.suppress(Exception):
                child.kill()
        if kernel.get("client"):
            kernel["client"].stop_channels()

    async def restart(self, params):
        kernel = self.get_kernel(params)
        await self.interrupt(params)
        async with kernel["lock"]:
            self.status(kernel, "restarting")
            await self.stop_manager(kernel)
            try:
                await self.launch(kernel)
                self.status(kernel, "idle")
            except Exception:
                self.status(kernel, "dead")
                raise
        return self.public_kernel(kernel)

    async def close(self, params):
        kernel = self.get_kernel(params)
        kernel["closing"] = True
        kernel["interrupted"] = True
        # Shutdown immediately also handles native extensions that ignore interrupts.
        await self.stop_manager(kernel)
        if kernel.get("collector"):
            kernel["collector"].cancel()
        self.kernels.pop(kernel["id"], None)
        self.status(kernel, "closed")
        return {"ok": True, "kernelId": kernel["id"]}

    async def list(self, params=None):
        return [self.public_kernel(k) for k in self.kernels.values()]

    async def shutdown(self, params=None):
        self.stopping = True
        await asyncio.gather(*(self.close({"kernelId": kid}) for kid in list(self.kernels)), return_exceptions=True)
        return {"ok": True}

    async def maintenance(self):
        try:
            import psutil
        except ImportError:
            psutil = None
        while not self.stopping:
            await asyncio.sleep(1)
            for kernel in list(self.kernels.values()):
                if kernel["status"] == "idle" and time.monotonic() - kernel["lastMonotonic"] > self.idle_seconds:
                    await self.close({"kernelId": kernel["id"]})
                    continue
                if psutil and kernel.get("manager"):
                    with contextlib.suppress(Exception):
                        pid = kernel["manager"].provisioner.pid
                        process = kernel.setdefault("process", psutil.Process(pid))
                        family = [process, *process.children(recursive=True)]
                        process_cache = kernel.setdefault("processCache", {})
                        processes = [process_cache.setdefault(item.pid, item) for item in family]
                        kernel["resources"] = {"pid": pid, "memoryBytes": sum(p.memory_info().rss for p in family),
                                                "cpuPercent": sum(p.cpu_percent() for p in processes), "sampledAt": now()}
                        self.event("status", kernelId=kernel["id"], status=kernel["status"], resources=kernel["resources"])
                        if kernel["status"] == "busy" and kernel["resources"]["memoryBytes"] > kernel["limits"]["maxMemoryBytes"]:
                            kernel["resourceLimit"] = f"Kernel exceeded its {kernel['limits']['maxMemoryBytes']} byte memory limit and was stopped. Restart it to continue."
                            self.status(kernel, "dead", kernel["activeExecutionId"])
                            await self.stop_manager(kernel)
                            if kernel.get("collector"):
                                kernel["collector"].cancel()

    async def run_process(self, argv, operation_id=None, timeout=300, env=None):
        process = await asyncio.create_subprocess_exec(*argv, stdin=asyncio.subprocess.DEVNULL, stdout=asyncio.subprocess.PIPE,
                                                       stderr=asyncio.subprocess.STDOUT, env=env or clean_env())
        output = bytearray()
        async def read():
            while chunk := await process.stdout.read(8192):
                if len(output) < 256 * 1024:
                    output.extend(chunk[:256 * 1024 - len(output)])
                    if operation_id:
                        self.event("environment-log", operationId=operation_id, text=chunk.decode("utf-8", errors="replace"))
            return await process.wait()
        try:
            code = await asyncio.wait_for(read(), timeout)
        except BaseException:
            descendants = []
            with contextlib.suppress(Exception):
                import psutil
                descendants = psutil.Process(process.pid).children(recursive=True)
            process.kill()
            for child in reversed(descendants):
                with contextlib.suppress(Exception):
                    child.kill()
            await process.wait()
            raise
        text = output.decode("utf-8", errors="replace")
        if code != 0:
            raise RuntimeError(f"Command failed ({code}): {text[-8000:]}")
        return text

    async def environments_list(self, params=None):
        result = [self.environment()]
        if self.local_r()["available"]:
            result.append(self.environment("r-default"))
        if self.environment_root.is_dir():
            for folder in self.environment_root.iterdir():
                with contextlib.suppress(ValueError, OSError, json.JSONDecodeError):
                    result.append(self.environment(folder.name))
        return result

    async def environments_create(self, params):
        environment_id = str(uuid.uuid4())
        folder = self.environment_root / environment_id
        self.environment_root.mkdir(parents=True, exist_ok=True)
        language = params.get("language", "python")
        if language not in ("python", "r"):
            raise ValueError("Environment language must be python or r")
        if language == "python":
            await self.run_process([sys.executable, "-m", "venv", str(folder)], environment_id)
        else:
            if not self.local_r()["available"]:
                raise ValueError("Install app-local R before creating an R environment")
            folder.mkdir()
        metadata = {"name": str(params.get("name", "Research environment"))[:100], "language": language, "createdAt": now()}
        (folder / "science-environment.json").write_text(json.dumps(metadata), encoding="utf-8")
        try:
            base_packages = ["ipykernel", "matplotlib", "numpy", "pandas"] if language == "python" else ["IRkernel", "jsonlite"]
            await self.environments_install({"environmentId": environment_id,
                "packages": [*base_packages, *(params.get("packages") or [])]})
        except Exception as exc:
            metadata["setupError"] = str(exc)
            (folder / "science-environment.json").write_text(json.dumps(metadata), encoding="utf-8")
            raise
        return self.environment(environment_id)

    async def environments_install(self, params):
        environment = self.environment(params.get("environmentId", "default"))
        if not environment["managed"]:
            raise ValueError("Package changes are restricted to app-managed environments")
        packages = params.get("packages")
        if not isinstance(packages, list) or not 1 <= len(packages) <= 50:
            raise ValueError("Provide 1 to 50 package requirements")
        if any(not isinstance(p, str) or len(p) > 200 or not PACKAGE_SPEC.fullmatch(p) for p in packages):
            raise ValueError("Use named PyPI packages with optional version constraints; URLs, paths and flags are not accepted")
        if any(k["environmentId"] == environment["id"] and k["status"] == "busy" for k in self.kernels.values()):
            raise ValueError("Stop active executions before changing their environment")
        lock = self.environment_locks.setdefault(environment["id"], asyncio.Lock())
        async with lock:
            if environment.get("language") == "r":
                if any(not re.fullmatch(r"[A-Za-z][A-Za-z0-9.]*", p) for p in packages):
                    raise ValueError("R package installs require CRAN package names without version constraints")
                expr = (f"pkgs <- c({','.join(json.dumps(p) for p in packages)}); "
                        "lib <- Sys.getenv('R_LIBS_USER'); "
                        "install.packages(pkgs, lib=lib, repos='https://cloud.r-project.org', type='binary'); "
                        "missing <- setdiff(pkgs, installed.packages(lib.loc=lib)[,'Package']); "
                        "if (length(missing)) stop(paste('Packages were not installed:', paste(missing, collapse=', ')))")
                logs = await self.run_process([self.local_r()["rscript"], "--vanilla", "-e", expr], environment["id"],
                                              timeout=600, env=self.r_env(environment["path"]))
            else:
                uv = shutil.which("uv")
                argv = [uv, "pip", "install", "--python", environment["pythonPath"]] if uv else [environment["pythonPath"], "-m", "pip", "install"]
                if environment["id"] == "default" and Path(environment["path"]).resolve().is_relative_to(self.runtime_root / "python"):
                    # The bundled uv standalone interpreter is marked externally
                    # managed. This exception only targets our isolated app copy.
                    argv.append("--break-system-packages")
                # Require wheels: no implicit build scripts during the package step.
                logs = await self.run_process([*argv, "--only-binary=:all:", *packages], environment["id"], timeout=600)
            return {"ok": True, "environmentId": environment["id"], "logs": logs,
                    "restartRequired": any(k["environmentId"] == environment["id"] for k in self.kernels.values())}

    async def environments_export(self, params):
        environment = self.environment(params.get("environmentId", "default"))
        if environment.get("language") == "r":
            expr = "p <- installed.packages(); cat(paste(p[,'Package'], p[,'Version'], sep='==', collapse='\\n'))"
            requirements = await self.run_process([self.local_r()["rscript"], "--vanilla", "-e", expr], timeout=30,
                                                 env=self.r_env(environment["path"]))
            return {"environmentId": environment["id"], "language": "r", "requirements": requirements, "exportedAt": now()}
        code = "import importlib.metadata as m; print('\\n'.join(sorted(d.metadata['Name']+'=='+d.version for d in m.distributions() if d.metadata['Name'])))"
        requirements = await self.run_process([environment["pythonPath"], "-c", code], timeout=30)
        return {"environmentId": environment["id"], "requirements": requirements, "exportedAt": now()}

    async def preview(self, params):
        path = self.contained(params.get("path", ""))
        if not path.is_file():
            raise ValueError("Preview path must be a file")
        rows = max(1, min(200, int(params.get("rows", 100))))
        suffix = path.suffix.lower()
        def read():
            import pandas as pd
            if suffix == ".parquet":
                import pyarrow.parquet as pq
                parquet = pq.ParquetFile(path)
                iterator = parquet.iter_batches(batch_size=rows)
                batch = next(iterator, None)
                frame = batch.to_pandas() if batch is not None else pd.DataFrame(columns=parquet.schema.names)
                total = parquet.metadata.num_rows
            elif suffix in (".xlsx", ".xlsm"):
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                try:
                    sheet = workbook.active
                    iterator = sheet.iter_rows(values_only=True)
                    headers = next(iterator, [])
                    records = []
                    for _, record in zip(range(rows), iterator):
                        records.append(record)
                    frame = pd.DataFrame(records, columns=[str(v) if v is not None else f"Column {i+1}" for i, v in enumerate(headers)])
                    total = max(0, (sheet.max_row or 1) - 1)
                finally:
                    workbook.close()
            elif suffix in (".csv", ".tsv"):
                frame = pd.read_csv(path, nrows=rows + 1, sep="\t" if suffix == ".tsv" else ",")
                total = None
            elif suffix == ".docx":
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(path) as archive:
                    info = archive.getinfo("word/document.xml")
                    if info.file_size > 8 * 1024 * 1024:
                        raise ValueError("Document text exceeds preview limit")
                    xml = archive.read(info)
                    root = ET.fromstring(xml)
                    text = "\n".join("".join(p.itertext()) for p in root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"))
                    return {"text": text[:100000], "truncated": len(text) > 100000, "size": path.stat().st_size}
            else:
                raise ValueError("Python preview supports parquet, xlsx, xlsm, csv, tsv and docx")
            record_json = frame.head(rows).to_json(orient="records", date_format="iso", default_handler=str)
            return {"columns": [str(v) for v in frame.columns], "rows": json.loads(record_json),
                    "totalRows": total, "truncated": total > rows if total is not None else len(frame) > rows,
                    "size": path.stat().st_size}
        return await asyncio.to_thread(read)

    async def dispatch(self, message):
        request_id = message.get("id")
        method = message.get("method", "").replace("/", "_").replace(".", "_")
        allowed = {"probe", "create", "execute", "interrupt", "restart", "close", "list", "shutdown",
                   "environments_list", "environments_create", "environments_install", "environments_export", "preview"}
        try:
            if method not in allowed:
                raise ValueError("Unknown worker method")
            result = await getattr(self, method)(message.get("params") or {})
            emit({"id": request_id, "result": result})
        except Exception as exc:
            emit({"id": request_id, "error": {"message": str(exc) or f"{type(exc).__name__} while running {method}", "type": type(exc).__name__}})


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--idle-seconds", type=int, default=1800)
    args = parser.parse_args()
    bridge = Bridge(args.root, max(10, args.idle_seconds))
    queue = asyncio.Queue()
    loop = asyncio.get_running_loop()
    def read_stdin():
        for line in sys.stdin:
            loop.call_soon_threadsafe(queue.put_nowait, line)
        with contextlib.suppress(RuntimeError):
            loop.call_soon_threadsafe(queue.put_nowait, None)
    threading.Thread(target=read_stdin, daemon=True).start()
    maintenance = asyncio.create_task(bridge.maintenance())
    try:
        while not bridge.stopping:
            line = await queue.get()
            if line is None:
                break
            try:
                if len(line) > 2 * 1024 * 1024:
                    raise ValueError("Request exceeds 2 MiB")
                message = json.loads(line)
                if not isinstance(message, dict):
                    raise ValueError("Request must be an object")
            except (ValueError, json.JSONDecodeError) as exc:
                emit({"id": None, "error": {"message": str(exc)}})
                continue
            task = asyncio.create_task(bridge.dispatch(message))
            bridge.tasks.add(task)
            task.add_done_callback(bridge.tasks.discard)
            if message.get("method") == "shutdown":
                await task
                break
    finally:
        maintenance.cancel()
        await bridge.shutdown()
        for task in list(bridge.tasks):
            task.cancel()
        await asyncio.gather(*bridge.tasks, return_exceptions=True)
        with contextlib.suppress(asyncio.CancelledError):
            await maintenance


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stdin.reconfigure(encoding="utf-8")
    asyncio.run(main())
